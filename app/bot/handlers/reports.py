"""Итоги смены — текстовый отчёт и Excel-выгрузка."""
import io
from datetime import date

from aiogram import Router
from aiogram.filters import Command
from aiogram.types import BufferedInputFile, Message

from app.db.base import async_session_factory
from app.db.crud import (
    get_active_session,
    get_all_active_staff,
    get_incidents_all,
    get_session_stats,
    get_staff_task_stats,
)
from app.db.models import STATUS_LABELS, Staff, StaffRole, TaskStatus

router = Router(name="reports")
_ADMIN_ROLES = {StaffRole.admin, StaffRole.senior_counselor}


# ── /report_session — текстовая сводка ───────────────────────────────────────

@router.message(Command("report_session"))
async def cmd_report_session(message: Message, staff: Staff | None = None):
    if staff is None or staff.role not in _ADMIN_ROLES:
        return
    async with async_session_factory() as session:
        sess = await get_active_session(session)
        if not sess:
            await message.answer("❌ Нет активной смены.")
            return
        stats = await get_session_stats(session, sess.id)
        all_staff = await get_all_active_staff(session)

    done = stats["done_tasks"]
    total = stats["total_tasks"]
    pct = round(done / total * 100) if total else 0

    lines = [
        f"📊 <b>Итоги смены «{sess.name}»</b>",
        f"Дата отчёта: {date.today().strftime('%d.%m.%Y')}\n",
        f"📋 Задачи: {done}/{total} выполнено ({pct}%)",
        f"🎉 Мероприятий: {stats['total_events']}",
        f"🚨 Инцидентов: {stats['total_incidents']}",
        f"👥 Персонал: {len(all_staff)} чел.",
    ]
    await message.answer("\n".join(lines), parse_mode="HTML")


# ── /report_excel — Excel-выгрузка ────────────────────────────────────────────

@router.message(Command("report_excel"))
async def cmd_report_excel(message: Message, staff: Staff | None = None):
    if staff is None or staff.role not in _ADMIN_ROLES:
        return
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        await message.answer("❌ openpyxl не установлен.")
        return

    async with async_session_factory() as session:
        sess = await get_active_session(session)
        if not sess:
            await message.answer("❌ Нет активной смены.")
            return
        stats = await get_session_stats(session, sess.id)
        all_staff = await get_all_active_staff(session)
        incidents = await get_incidents_all(session, session_id=sess.id, limit=500)

        staff_stats = {}
        for s in all_staff:
            staff_stats[s] = await get_staff_task_stats(session, s.id, sess.id)

    wb = openpyxl.Workbook()
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font_white = Font(bold=True, color="FFFFFF")

    def _header(ws, cols):
        ws.append(cols)
        for cell in ws[1]:
            cell.font = header_font_white
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

    # Sheet 1: Summary
    ws = wb.active
    ws.title = "Сводка"
    ws.append(["Смена", sess.name])
    ws.append(["Дата отчёта", date.today().strftime("%d.%m.%Y")])
    ws.append(["Задач выполнено", f"{stats['done_tasks']}/{stats['total_tasks']}"])
    ws.append(["Мероприятий", stats["total_events"]])
    ws.append(["Инцидентов", stats["total_incidents"]])
    ws.append(["Персонал", len(all_staff)])
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 30

    # Sheet 2: Staff stats
    ws2 = wb.create_sheet("Персонал")
    _header(ws2, ["Сотрудник", "Роль", "Задач всего", "Выполнено", "Просрочено"])
    for s, s_stats in staff_stats.items():
        total_t = sum(s_stats.values())
        done_t = s_stats.get(TaskStatus.done.value, 0)
        overdue_t = s_stats.get(TaskStatus.overdue.value, 0)
        ws2.append([s.full_name, s.role.value, total_t, done_t, overdue_t])
    for col in ["A", "B", "C", "D", "E"]:
        ws2.column_dimensions[col].width = 20

    # Sheet 3: Incidents
    ws3 = wb.create_sheet("Инциденты")
    _header(ws3, ["Дата/время", "Тип", "Репортер", "Описание", "Ребёнок"])
    for inc in incidents:
        child_name = inc.child.full_name if inc.child else "—"
        ws3.append([
            inc.created_at.strftime("%d.%m.%Y %H:%M"),
            inc.type,
            inc.reporter.full_name,
            inc.description,
            child_name,
        ])
    for col, width in [("A", 18), ("B", 18), ("C", 22), ("D", 40), ("E", 22)]:
        ws3.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"report_{sess.name.replace(' ', '_')}_{date.today()}.xlsx"
    await message.answer_document(
        BufferedInputFile(buf.read(), filename=filename),
        caption=f"📊 Отчёт по смене «{sess.name}»",
    )
